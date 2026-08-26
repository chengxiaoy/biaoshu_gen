import json
from pathlib import Path

import pytest
from docx import Document
from pydantic_ai import Agent
from pydantic_ai.messages import ModelResponse, ToolCallPart
from pydantic_ai.models.function import AgentInfo, FunctionModel

from biaoshu_gen.nodes import extract_template as et
from biaoshu_gen.schemas import TemplateAnchor
from biaoshu_gen.state import BidState, run_dir


def _tender(tmp_path: Path) -> Path:
    """合成招标文件:非格式章节在前,第七章格式章节(段落+表格)在后。

    前导空段使 body 子元素下标与块编号错开(空段不编号但占 element_index),
    确保「blocks[i].element_index 映射」不被「直接拿 i 当元素下标」的退化实现蒙混。
    """
    p = tmp_path / "tender.docx"
    doc = Document()
    doc.add_paragraph("")                                            # 空段:块 0 实为元素 1
    doc.add_paragraph("第二章 投标人须知")
    doc.add_paragraph("递交截止时间为开标之日。")
    doc.add_paragraph("第七章 投标文件的格式")
    doc.add_paragraph("投标函（格式）")
    doc.add_paragraph("兹承诺按招标文件要求投标。")
    t = doc.add_table(rows=1, cols=1)
    t.cell(0, 0).text = "报价表"
    doc.save(p)
    return p


def _fake_make(responses: list[dict]):
    """按调用次序返回预设 JSON 的假 agent 工厂;记录收到的 prompt(structure 测试同款)。"""
    calls: list[str] = []

    def make(output_type, system_prompt, retries=2):
        def fn(messages, info: AgentInfo):
            calls.append(messages[-1].parts[-1].content)
            tool = info.output_tools[0].name if info.output_tools else "final_result"
            payload = responses[min(len(calls) - 1, len(responses) - 1)]
            return ModelResponse(parts=[ToolCallPart(tool_name=tool, args=json.dumps(payload))])
        return Agent(model=FunctionModel(fn), output_type=output_type,
                     system_prompt=system_prompt, retries=retries)

    make.calls = calls
    return make


def test_validate_anchor_normalizes_and_rejects():
    assert et._validate_anchor(TemplateAnchor(start_index=2), 6) == (2, None)
    assert et._validate_anchor(TemplateAnchor(start_index=2, end_index=6), 6) == (2, None)  # end==n 视同文末
    assert et._validate_anchor(TemplateAnchor(start_index=2, end_index=4), 6) == (2, 4)
    with pytest.raises(ValueError):
        et._validate_anchor(TemplateAnchor(start_index=6), 6)          # start 越界
    with pytest.raises(ValueError):
        et._validate_anchor(TemplateAnchor(start_index=3, end_index=3), 6)  # end 不严格大于 start


def test_node_extracts_via_llm_bounds(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    tender = _tender(tmp_path)
    state = BidState(run_id="run-1", tender_path=str(tender))
    make = _fake_make([{"start_index": 2, "end_index": None}])      # 块2 = 第七章标题
    monkeypatch.setattr(et, "make_agent", make)

    updates = et.extract_template_node(state)
    ws = run_dir(state) / "02_template"
    tpl = ws / "标书模板.docx"
    assert updates["template_docx_path"] == str(tpl)

    texts = [p.text for p in Document(str(tpl)).paragraphs]
    # 前导空段使块下标(2)≠元素下标(3):错位剪裁会混入上一段「递交截止时间…」,在此被钉死
    assert texts == ["第七章 投标文件的格式", "投标函（格式）", "兹承诺按招标文件要求投标。"]
    tpl_md = (ws / "template.md").read_text(encoding="utf-8")
    # 合成样本无 Heading 样式 -> derive_template_md 走扁平列表兜底(标题+表格 stub 均应出现)
    assert "第七章 投标文件的格式" in tpl_md and "【表格】" in tpl_md
    assert (ws / "report.md").exists()
    assert "[2] 第七章 投标文件的格式" in make.calls[0]              # 块化行进 prompt


def test_node_retries_once_with_error_feedback(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    tender = _tender(tmp_path)
    state = BidState(run_id="run-1", tender_path=str(tender))
    make = _fake_make([{"start_index": 99}, {"start_index": 2}])
    monkeypatch.setattr(et, "make_agent", make)

    et.extract_template_node(state)
    assert len(make.calls) == 2
    assert "错误" in make.calls[1]                                   # 第二次 prompt 带错误反馈


def test_node_raises_after_retry_exhausted(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    tender = _tender(tmp_path)
    state = BidState(run_id="run-1", tender_path=str(tender))
    make = _fake_make([{"start_index": 99}])
    monkeypatch.setattr(et, "make_agent", make)

    with pytest.raises(et.TemplateExtractError):
        et.extract_template_node(state)
    assert len(make.calls) == 2


def test_node_sidecar_copies_directly_without_llm(tmp_path, monkeypatch):
    """随附权威模板直通复制为底稿,不走 LLM 定界。"""
    monkeypatch.chdir(tmp_path)
    tender = _tender(tmp_path)
    sidecar = tmp_path / "投标模板.docx"
    sd = Document()
    sd.add_paragraph("随附模板正文")
    sd.save(sidecar)
    state = BidState(run_id="run-1", tender_path=str(tender),
                     template_docx_path=str(sidecar))

    def boom(*a, **k):
        raise AssertionError("随附模板存在时不应调用 LLM")
    monkeypatch.setattr(et, "make_agent", boom)

    updates = et.extract_template_node(state)
    ws = run_dir(state) / "02_template"
    assert [p.text for p in Document(str(ws / "标书模板.docx")).paragraphs] == ["随附模板正文"]
    assert (ws / "template.md").exists() and (ws / "report.md").exists()
    assert updates["template_docx_path"] == str(ws / "标书模板.docx")
