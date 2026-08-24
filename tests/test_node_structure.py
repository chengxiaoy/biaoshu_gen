from biaoshu_gen.docx_io import DocxSection, NumberedBlock
from biaoshu_gen.nodes import structure as st
from biaoshu_gen.schemas import StructureHeading, StructureOutline


def _block(i: int, text: str) -> NumberedBlock:
    return NumberedBlock(index=i, kind="p", stub=text, md=text)


def test_validate_headings_drops_bad_and_clamps():
    outline = StructureOutline.model_validate({"headings": [
        {"index": 3, "level": 1, "title": "第三章 评标办法"},
        {"index": 1, "level": 1, "title": "第一章 总体要求"},      # 乱序 -> 丢弃
        {"index": 5, "level": 9, "title": "第五章 附则"},          # level 夹取 3
        {"index": 7, "level": 2, "title": ""},                     # 空 title 丢弃
        {"index": 8, "level": 2, "title": "x" * 51},               # 超 50 字丢弃
        {"index": 9, "level": 0, "title": "第六章 其他"},          # level 夹取 1
    ]})
    hs = st.validate_headings(outline, n_blocks=12)
    assert [(h.index, h.level, h.title) for h in hs] == [
        (3, 1, "第三章 评标办法"), (5, 3, "第五章 附则"), (9, 1, "第六章 其他")]
    assert hs[0].index == 3                                        # 幸存者首位强制 level 1 已是 1


def test_split_by_headings_assigns_content_and_preamble():
    blocks = [_block(i, t) for i, t in enumerate([
        "封面文字", "第一章 总体要求", "系统需支持 1000 并发。",
        "1.1 性能指标", "响应时间 ≤ 2 秒。", "第二章 商务条款", "质保三年。",
    ])]
    headings = [StructureHeading(index=1, level=1, title="第一章 总体要求"),
                StructureHeading(index=3, level=3, title="1.1 性能指标"),
                StructureHeading(index=5, level=1, title="第二章 商务条款")]
    secs = st.split_by_headings(blocks, headings)
    assert [(s.level, s.title) for s in secs] == [
        (0, "(前言)"), (1, "第一章 总体要求"), (3, "1.1 性能指标"), (1, "第二章 商务条款")]
    assert secs[0].content == "封面文字"
    assert "1000 并发" in secs[1].content
    assert "响应时间" in secs[2].content
    assert "质保三年" in secs[3].content


def test_split_by_headings_drops_empty_preamble():
    blocks = [_block(i, t) for i, t in enumerate(["第一章 总则", "正文若干。"])]
    secs = st.split_by_headings(blocks, [StructureHeading(index=0, level=1, title="第一章 总则")])
    assert [(s.level, s.title) for s in secs] == [(1, "第一章 总则")]


import json

import pytest
from docx import Document
from pydantic_ai import Agent
from pydantic_ai.messages import ModelResponse, ToolCallPart
from pydantic_ai.models.function import AgentInfo, FunctionModel

from biaoshu_gen.schemas import StructureOutline


def _unstructured_tender(tmp_path):
    p = tmp_path / "ns.docx"
    doc = Document()
    doc.add_paragraph("第一章 采购需求")
    doc.add_paragraph("内容甲。" * 30)
    doc.add_paragraph("第二章 评标办法")
    doc.add_paragraph("内容乙。" * 30)
    doc.save(p)
    return p


def _fake_make(responses: list[str]):
    """按调用次序返回预设 JSON 的假 agent 工厂;记录收到的 prompt。"""
    calls: list[str] = []

    def make(output_type, system_prompt, retries=2):
        def fn(messages, info: AgentInfo):
            calls.append(messages[-1].parts[-1].content)
            tool = info.output_tools[0].name if info.output_tools else "final_result"
            payload = responses[min(len(calls) - 1, len(responses) - 1)]
            return ModelResponse(parts=[ToolCallPart(tool_name=tool, args=payload)])
        return Agent(model=FunctionModel(fn), output_type=output_type,
                     system_prompt=system_prompt, retries=retries)

    make.calls = calls
    return make


def test_rebuild_sections_happy_path(tmp_path, monkeypatch):
    p = _unstructured_tender(tmp_path)
    make = _fake_make([json.dumps({"headings": [
        {"index": 0, "level": 1, "title": "第一章 采购需求"},
        {"index": 2, "level": 1, "title": "第二章 评标办法"},
    ]}, ensure_ascii=False)])
    monkeypatch.setattr(st, "make_agent", make)
    secs = st.rebuild_sections(p)
    assert [(s.level, s.title) for s in secs] == [
        (1, "第一章 采购需求"), (1, "第二章 评标办法")]
    assert "内容甲" in secs[0].content
    prompt = make.calls[0]
    assert "[0] 第一章 采购需求" in prompt and "【表格】" not in prompt
    assert "目录" in prompt                                    # 去目录指引在 prompt 中


def test_rebuild_sections_retries_once_with_error_feedback(tmp_path, monkeypatch):
    p = _unstructured_tender(tmp_path)
    # 首次输出必须真正未过校验才触发重试:index 合法但 title 为空 -> validate_headings 丢弃
    bad = json.dumps({"headings": [{"index": 2, "level": 1, "title": ""}]})
    good = json.dumps({"headings": [
        {"index": 0, "level": 1, "title": "第一章 采购需求"},
        {"index": 2, "level": 1, "title": "第二章 评标办法"}]}, ensure_ascii=False)
    make = _fake_make([bad, good])
    monkeypatch.setattr(st, "make_agent", make)
    secs = st.rebuild_sections(p)
    assert len(make.calls) == 2
    assert "错误" in make.calls[1]                             # 第二次 prompt 带错误反馈
    assert len(secs) == 2


def test_rebuild_sections_raises_after_retry_exhausted(tmp_path, monkeypatch):
    p = _unstructured_tender(tmp_path)
    bad = json.dumps({"headings": [{"index": 5, "level": 1, "title": "越界"}]})
    make = _fake_make([bad])
    monkeypatch.setattr(st, "make_agent", make)
    with pytest.raises(st.StructureError):
        st.rebuild_sections(p)
    assert len(make.calls) == 2
