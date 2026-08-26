import json
from pathlib import Path

import pytest
from docx import Document
from pydantic_ai import Agent
from pydantic_ai.messages import ModelResponse, ToolCallPart
from pydantic_ai.models.function import AgentInfo, FunctionModel

from biaoshu_gen.docx_io import find_deviation_tables
from biaoshu_gen.nodes import deviation_table as dev
from biaoshu_gen.state import BidState, run_dir

_HDR = ["序号", "磋商文件章节条款号", "磋商文件要求", "响应文件的应答", "偏离说明"]


def _state(tmp_path: Path, monkeypatch) -> BidState:
    """含两块偏离表的响应模板 + parse/facts 产物。"""
    monkeypatch.chdir(tmp_path)
    tpl = tmp_path / "标书模板.docx"
    doc = Document()
    doc.add_paragraph("一、磋商响应声明")
    doc.add_paragraph("七、合同条款偏离表")
    t1 = doc.add_table(rows=2, cols=5)
    for i, h in enumerate(_HDR):
        t1.cell(0, i).text = h
    doc.add_paragraph("八、采购需求偏离表")
    t2 = doc.add_table(rows=3, cols=5)
    for i, h in enumerate(_HDR):
        t2.cell(0, i).text = h
    doc.save(tpl)

    state = BidState(run_id="run-1", tender_path=str(tmp_path / "tender.docx"),
                     template_docx_path=str(tpl))
    parse = run_dir(state) / "01_parse"
    parse.mkdir(parents=True)
    (parse / "requirements.yaml").write_text("tech_requirements:\n- 质保期3年\n", encoding="utf-8")
    (parse / "invalidation.yaml").write_text("items: []\n", encoding="utf-8")
    (run_dir(state) / "03_facts.yaml").write_text("schedule: 30天\n", encoding="utf-8")
    return state


def _fake_make(responses: list[dict]):
    """按调用次序返回预设 JSON 的假 agent 工厂;记录收到的 prompt。"""
    calls: list[str] = []

    def make(output_type, system_prompt, retries=2):
        def fn(messages, info: AgentInfo):
            calls.append(messages[-1].parts[-1].content)
            tool = info.output_tools[0].name if info.output_tools else "final_result"
            payload = responses[min(len(calls) - 1, len(responses) - 1)]
            return ModelResponse(parts=[ToolCallPart(tool_name=tool, args=json.dumps(payload))])
        return Agent(model=FunctionModel(fn), output_type=output_type,
                     system_prompt=system_prompt, retries=retries)

    make.calls = calls
    return make


_ROWS = {
    "contract_rows": [{"clause": "第12条", "requirement": "交货期30天",
                       "response": "承诺30天交货", "deviation": "无偏离"}],
    "requirement_rows": [{"clause": "3.2", "requirement": "质保期3年",
                          "response": "满足,质保3年", "deviation": ""},
                         {"clause": "3.5", "requirement": "现场安装",
                          "response": "负责安装调试", "deviation": "正偏离"}],
}


def test_node_fills_both_tables_via_llm(tmp_path, monkeypatch):
    state = _state(tmp_path, monkeypatch)
    make = _fake_make([_ROWS])
    monkeypatch.setattr(dev, "make_agent", make)

    updates = dev.deviation_table_node(state)
    out = Path(updates["deviation_docx_path"])
    assert out.name == "deviation.docx" and out.exists()

    doc = Document(str(out))
    tables = {kind: t for t, kind in find_deviation_tables(doc)}
    assert [c.text for c in tables["contract"].rows[1].cells] == \
        ["1", "第12条", "交货期30天", "承诺30天交货", "无偏离"]
    assert len(tables["requirement"].rows) == 3                      # 表头+2行,旧空行已清
    assert tables["requirement"].rows[1].cells[0].text == "1"        # 序号代码生成
    assert tables["requirement"].rows[2].cells[4].text == "正偏离"
    assert [c.text for c in tables["requirement"].rows[1].cells][2] == "质保期3年"

    prompt = make.calls[0]
    assert "磋商文件章节条款号" in prompt                             # 表头进 prompt
    assert "质保期3年" in prompt and "30天" in prompt                  # requirements/facts 进 prompt
    assert "采购需求偏离表" in prompt                                  # 表归类可辨


def test_node_retries_once_on_invalid_rows(tmp_path, monkeypatch):
    state = _state(tmp_path, monkeypatch)
    make = _fake_make([{"contract_rows": [], "requirement_rows": []}, _ROWS])
    monkeypatch.setattr(dev, "make_agent", make)

    dev.deviation_table_node(state)
    assert len(make.calls) == 2
    assert "校验" in make.calls[1]                                    # 第二次带错误反馈


def test_node_raises_after_retry_exhausted(tmp_path, monkeypatch):
    state = _state(tmp_path, monkeypatch)
    make = _fake_make([{"contract_rows": [], "requirement_rows": []}])
    monkeypatch.setattr(dev, "make_agent", make)

    with pytest.raises(dev.DeviationFillError):
        dev.deviation_table_node(state)
    assert len(make.calls) == 2


def test_node_rejects_rows_for_absent_table_kind(tmp_path, monkeypatch):
    """模板只有采购需求偏离表时,contract_rows 非空必须校验失败并带错重试。"""
    monkeypatch.chdir(tmp_path)
    tpl = tmp_path / "标书模板.docx"
    doc = Document()
    doc.add_paragraph("八、采购需求偏离表")
    t = doc.add_table(rows=2, cols=5)
    for i, h in enumerate(_HDR):
        t.cell(0, i).text = h
    doc.save(tpl)
    state = BidState(run_id="run-1", tender_path=str(tpl), template_docx_path=str(tpl))
    valid = {"contract_rows": [], "requirement_rows": _ROWS["requirement_rows"]}
    make = _fake_make([_ROWS, valid])                  # 第一次带 contract 行(模板无此表)
    monkeypatch.setattr(dev, "make_agent", make)

    updates = dev.deviation_table_node(state)
    assert len(make.calls) == 2
    assert "无contract类偏离表" in make.calls[1]        # 错误反馈点名缺失类别
    out = Document(updates["deviation_docx_path"])
    t2 = find_deviation_tables(out)[0][0]
    assert len(t2.rows) == 3                            # requirement 行已写入


def test_node_skips_without_deviation_table(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    tpl = tmp_path / "标书模板.docx"
    doc = Document()
    doc.add_paragraph("一、磋商响应声明")
    doc.save(tpl)
    state = BidState(run_id="run-1", tender_path=str(tpl), template_docx_path=str(tpl))

    def boom(*a, **k):
        raise AssertionError("无偏离表时不应调用 LLM")
    monkeypatch.setattr(dev, "make_agent", boom)

    assert dev.deviation_table_node(state) == {"deviation_docx_path": ""}
