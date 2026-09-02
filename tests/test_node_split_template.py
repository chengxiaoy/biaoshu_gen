import json
from pathlib import Path

import pytest
from docx import Document
from pydantic_ai import Agent
from pydantic_ai.messages import ModelResponse, ToolCallPart
from pydantic_ai.models.function import AgentInfo, FunctionModel

from biaoshu_gen.nodes import split_template as st
from biaoshu_gen.state import BidState, run_dir

_SECTIONS = [
    "一、磋商响应声明", "二、供应商资格证明文件", "三、保证金（本项目无须提供）",
    "四、报价表及分项价格表", "五、货物说明一览表", "六、项目实施方案",
    "七、合同条款偏离表", "八、采购需求偏离表", "九、类似业绩",
]


def _titled_tpl(path: Path) -> None:
    """带 Heading 样式的响应模板:前言 + 九个二级节,偏离节下带表。"""
    doc = Document()
    doc.add_paragraph("第五章 响应文件组成")
    for title in _SECTIONS:
        doc.add_heading(title, level=2)
        if "偏离表" in title:
            t = doc.add_table(rows=2, cols=2)
            t.cell(0, 0).text = "序号"
            t.cell(0, 1).text = "偏离说明"
        else:
            doc.add_paragraph(f"{title} 正文。")
    doc.save(path)


def _state(tmp_path: Path, monkeypatch) -> BidState:
    monkeypatch.chdir(tmp_path)
    tpl = tmp_path / "tpl.docx"
    _titled_tpl(tpl)
    ws = tmp_path / "run-1" / "02_template"
    ws.mkdir(parents=True)
    tpl_final = ws / "标书模板.docx"
    tpl_final.write_bytes(tpl.read_bytes())
    return BidState(run_id="run-1", tender_path=str(tpl),
                    template_docx_path=str(tpl_final))


def test_split_by_heading_rules(tmp_path, monkeypatch):
    state = _state(tmp_path, monkeypatch)

    def boom(*a, **k):
        raise AssertionError("有标题模板不应调用 LLM")
    monkeypatch.setattr(st, "make_agent", boom)

    updates = st.split_template_node(state)
    parts = updates["template_parts"]
    assert set(parts) == {"deviation", "technical", "forms", "commercial"}

    def texts(path_key: str) -> list[str]:
        return [p.text for p in Document(path_key).paragraphs if p.text.strip()]

    manifest = st.read_parts_yaml(run_dir(state))
    entries = {e["key"]: e for e in manifest["entries"]}
    # run 粒度:同桶不相邻区间各自成文件;主键文件只含首段
    assert texts(parts["forms"]) == ["一、磋商响应声明", "一、磋商响应声明 正文。",
                                     "二、供应商资格证明文件", "二、供应商资格证明文件 正文。"]
    assert [t for t in texts(entries["forms_2"]["path"])
            if not t.startswith("第五章")] == \
        ["四、报价表及分项价格表", "四、报价表及分项价格表 正文。",
         "五、货物说明一览表", "五、货物说明一览表 正文。"]
    assert texts(parts["commercial"]) == ["第五章 响应文件组成"]
    assert len(texts(entries["commercial_2"]["path"])) == 2            # 三、保证金 两段
    assert len(texts(entries["commercial_3"]["path"])) == 2            # 九、类似业绩
    assert texts("deviation" if False else parts["deviation"]) == \
        ["七、合同条款偏离表", "八、采购需求偏离表"]
    assert len(Document(parts["deviation"]).tables) == 2
    assert texts(parts["technical"]) == ["六、项目实施方案", "六、项目实施方案 正文。"]

    # parts.yaml:entries-only;primary 标记各桶首段,序列按文档原序
    assert [e["key"] for e in manifest["entries"]] == [
        "commercial", "forms", "commercial_2", "forms_2",
        "technical", "deviation", "commercial_3"]
    assert [e["primary"] for e in manifest["entries"]] == [
        True, True, False, False, True, True, False]


def test_split_untitled_template_via_llm(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    tpl = tmp_path / "tpl.docx"
    doc = Document()
    for text in ("一、磋商响应声明", "投标函正文。", "六、项目实施方案",
                 "方案正文。", "七、合同条款偏离表"):
        doc.add_paragraph(text)                     # 全普通段落,无 Heading 样式
    doc.save(tpl)
    ws = tmp_path / "run-1" / "02_template"
    ws.mkdir(parents=True)
    tpl_final = ws / "标书模板.docx"
    tpl_final.write_bytes(tpl.read_bytes())
    state = BidState(run_id="run-1", tender_path=str(tpl),
                     template_docx_path=str(tpl_final))

    calls: list[str] = []

    def make(output_type, system_prompt, retries=2):
        def fn(messages, info: AgentInfo):
            calls.append(messages[-1].parts[-1].content)
            tool = info.output_tools[0].name if info.output_tools else "final_result"
            payload = {"spans": [
                {"start_index": 0, "end_index": 2, "bucket": "forms"},
                {"start_index": 2, "end_index": 4, "bucket": "technical"},
                {"start_index": 4, "end_index": None, "bucket": "deviation"},
            ]}
            return ModelResponse(parts=[ToolCallPart(tool_name=tool, args=json.dumps(payload))])
        return Agent(model=FunctionModel(fn), output_type=output_type,
                     system_prompt=system_prompt, retries=retries)

    monkeypatch.setattr(st, "make_agent", make)
    updates = st.split_template_node(state)
    parts = updates["template_parts"]

    def texts(bucket: str) -> list[str]:
        return [p.text for p in Document(parts[bucket]).paragraphs if p.text.strip()]

    assert texts("forms") == ["一、磋商响应声明", "投标函正文。"]
    assert texts("technical") == ["六、项目实施方案", "方案正文。"]
    assert texts("deviation") == ["七、合同条款偏离表"]
    assert parts.get("commercial", "") == ""               # 空 bucket 不落文件
    assert len(calls) == 1                                 # 单次 LLM 调用


def test_split_llm_retry_on_bad_span(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    tpl = tmp_path / "tpl.docx"
    doc = Document()
    for text in ("一、磋商响应声明", "七、合同条款偏离表"):
        doc.add_paragraph(text)
    doc.save(tpl)
    ws = tmp_path / "run-1" / "02_template"
    ws.mkdir(parents=True)
    (ws / "标书模板.docx").write_bytes(tpl.read_bytes())
    state = BidState(run_id="run-1", tender_path=str(tpl), template_docx_path=str(ws / "标书模板.docx"))

    calls: list[str] = []

    def make(output_type, system_prompt, retries=2):
        def fn(messages, info: AgentInfo):
            calls.append(messages[-1].parts[-1].content)
            tool = info.output_tools[0].name if info.output_tools else "final_result"
            payload = ({"spans": [{"start_index": 99, "bucket": "forms"}]} if len(calls) == 1
                       else {"spans": [{"start_index": 0, "bucket": "forms"}]})
            return ModelResponse(parts=[ToolCallPart(tool_name=tool, args=json.dumps(payload))])
        return Agent(model=FunctionModel(fn), output_type=output_type,
                     system_prompt=system_prompt, retries=retries)

    monkeypatch.setattr(st, "make_agent", make)
    st.split_template_node(state)
    assert len(calls) == 2 and "校验" in calls[1]


def test_interleaved_bucket_yields_run_entries(tmp_path: Path, monkeypatch):
    """同桶多次出现(第七章框架内:投标函→(四)(五)商务→资格)须按连续区间落
    多份文件并登记 entries;order 仍按文档原序,legacy parts 只留各桶首段。"""
    monkeypatch.chdir(tmp_path)
    tpl = tmp_path / "标书模板.docx"
    d = Document()
    d.add_heading("第七章  投标文件的格式", level=1)     # commercial 头
    d.add_paragraph("头填充。")
    d.add_heading("投标函及报价文件", level=2)            # forms run A
    d.add_paragraph("函A。")
    d.add_heading("（四）法定代表人（负责人）身份证明", level=2)   # commercial 嵌入段!
    d.add_paragraph("身份证明体。")
    d.add_heading("（五）法定代表人（负责人）授权书", level=2)
    d.add_paragraph("授权书体。")
    d.add_heading("资格证明文件", level=2)                # forms run B(同桶第二段!)
    d.add_paragraph("资格体。")
    d.add_heading("技术部分", level=2)                    # technical
    d.add_paragraph("技术体。")
    d.save(tpl)

    state = BidState(run_id="run-1", tender_path=str(tpl), template_docx_path=str(tpl))
    updates = st.split_template_node(state)

    man = st.read_parts_yaml(run_dir(state))
    entries = man["entries"]
    seq = [(e["bucket"], e["first_element_index"]) for e in entries]
    # (四)(五)相邻归同一 run——与真实 software 模板((四)(五)嵌在投标函与资格间)一致
    assert seq == [("commercial", 0), ("forms", 2), ("commercial", 4),
                   ("forms", 8), ("technical", 10)]
    assert entries[2]["key"] == "commercial_2"
    assert entries[3]["key"] == "forms_2"
    assert entries[2]["path"] != entries[0]["path"]                     # 同桶第二段独立文件
    assert Path(updates["template_parts"]["commercial"]).exists()      # 主条目仍按旧键可用
    assert len(list((run_dir(state) / "02_template" / "parts").glob("*.docx"))) >= 5

    # 每份 run 文件只含自己区间的内容
    mid = Document(entries[2]["path"])
    texts = [p.text for p in mid.paragraphs if p.text.strip()]
    assert any("身份证明" in t or "授权书" in t for t in texts)
    assert not any("投标函" in t for t in texts)                        # 不含 forms 区内容
