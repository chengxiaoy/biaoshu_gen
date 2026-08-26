"""节点 2：投标模板抽取（非 harness）：LLM 定界 + python 剪裁 + 确定性派生 md。

从招标文件中定位"投标文件/响应文件的格式"章节（LLM 只定界），整包副本删区间
剪裁出 标书模板.docx；template.md / report.md 由副本确定性派生，零 LLM。
设计见 docs/superpowers/specs/2026-08-24-extract-template-non-harness-design.md。
"""
import shutil
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from ..docx_io import clip_docx, docx_to_sections, iter_numbered_blocks
from ..models import make_agent, run_sync       # noqa: F401  (测试 monkeypatch et.make_agent)
from ..prompts.extract_template import SYSTEM, build_user_prompt
from ..schemas import TemplateAnchor
from ..state import BidState, run_dir

_RETRY_TIMES = 2   # 首次 + 校验失败重试一次


class TemplateExtractError(RuntimeError):
    """模板抽取失败（LLM 两次输出均未通过校验，或剪裁结果为空）。"""


def _validate_anchor(anchor: TemplateAnchor, n_blocks: int) -> tuple[int, int | None]:
    """硬校验并归一化：(start, end|None) 块序号；end==n 视同 None（到文末）。"""
    s, e = anchor.start_index, anchor.end_index
    if not 0 <= s < n_blocks:
        raise ValueError(f"start_index {s} 越界（共 {n_blocks} 块）")
    if e is not None:
        if e == n_blocks:
            e = None
        elif not s < e < n_blocks:
            raise ValueError(f"end_index {e} 非法（需 start < end < n 或 == n）")
    return s, e


def _extract_from_tender(tender: Path, tpl_docx: Path) -> None:
    """块化 -> 单次 LLM 定界 -> 硬校验（失败带错重试一次）-> 映射元素下标剪裁。"""
    doc = Document(str(tender))
    blocks = iter_numbered_blocks(doc)
    if not blocks:
        raise TemplateExtractError(f"模板抽取失败：{tender} 无任何内容块")
    n_children = len(list(doc.element.body.iterchildren()))
    blocks_text = "\n".join(f"[{b.index}] {b.stub}" for b in blocks)
    agent = make_agent(TemplateAnchor, SYSTEM)

    prompt = build_user_prompt(blocks_text)
    rng: tuple[int, int | None] | None = None
    err = ""
    for _ in range(_RETRY_TIMES):
        anchor: TemplateAnchor = run_sync(agent, prompt).output
        try:
            rng = _validate_anchor(anchor, len(blocks))
            break
        except ValueError as exc:
            err = str(exc)
            prompt = build_user_prompt(blocks_text) + (
                f"\n\n上一次输出未通过校验（错误:{err}），请修正后重新输出。")
    if rng is None:
        raise TemplateExtractError(
            f"模板抽取失败：{tender} 两次输出均未通过校验（最后错误:{err}）")

    s_blk, e_blk = rng
    start_el = blocks[s_blk].element_index
    end_el = blocks[e_blk].element_index if e_blk is not None else n_children
    clip_docx(tender, tpl_docx, start_el, end_el)

    clipped = Document(str(tpl_docx))
    if not any(c.tag != qn("w:sectPr") for c in clipped.element.body.iterchildren()):
        raise TemplateExtractError(f"剪裁结果为空：{tender}（区间 [{start_el},{end_el})）")


def derive_template_md(tpl_docx: Path) -> str:
    """对模板副本确定性派生说明：标题树 + 表格类/文档类标注 + 字数（零 LLM）。"""
    secs = docx_to_sections(tpl_docx)
    titled = [s for s in secs if s.level]
    lines = ["# 响应文件模板说明", "",
             "> 本文件由 标书模板.docx 确定性派生：目录树 + 填写方式标注 + 字数。", ""]
    if not titled:   # 无标题样式退化：列内容块摘要，不报错
        lines += [f"- {b.stub}" for b in iter_numbered_blocks(Document(str(tpl_docx)))]
        return "\n".join(lines) + "\n"
    for s in titled:
        kind = "«表格类»" if "|" in s.content else "«文档类»"
        lines.append("  " * (s.level - 1) + f"- {s.title}（{len(s.content)}字）{kind}")
    return "\n".join(lines) + "\n"


def derive_report_md(tpl_docx: Path) -> str:
    """人读版报告：模板概览与各节字数/表格行一览（零 LLM）。"""
    import re

    sep = re.compile(r"^\|(?:\s*:?-+:?\s*\|)+$")
    secs = docx_to_sections(tpl_docx)
    titled = [s for s in secs if s.level]
    lines = ["# 响应文件模板抽取报告", "",
             f"- 模板副本：标书模板.docx（{tpl_docx.stat().st_size} 字节）",
             f"- 章节数：{len(titled)}；总字数：{sum(len(s.content) for s in secs)}",
             "", "| 章节 | 层级 | 字数 | 表格行 |", "|---|---|---|---|"]
    for s in titled:
        rows = sum(1 for ln in s.content.splitlines()
                   if ln.strip().startswith("|") and not sep.match(ln.strip()))
        lines.append(f"| {s.title} | H{s.level} | {len(s.content)} | {rows} |")
    return "\n".join(lines) + "\n"


def extract_template_node(state: BidState) -> dict:
    d = run_dir(state)
    ws = d / "02_template"
    ws.mkdir(parents=True, exist_ok=True)
    tpl_docx = ws / "标书模板.docx"

    attached = Path(state.template_docx_path) if state.template_docx_path else None
    if attached and attached.exists():
        shutil.copyfile(attached, tpl_docx)   # 随附权威模板直通，跳过定界与剪裁
    else:
        _extract_from_tender(Path(state.tender_path), tpl_docx)

    (ws / "template.md").write_text(derive_template_md(tpl_docx), encoding="utf-8")
    (ws / "report.md").write_text(derive_report_md(tpl_docx), encoding="utf-8")
    return {"template_docx_path": str(tpl_docx)}
